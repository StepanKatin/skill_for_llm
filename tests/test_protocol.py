"""Protocol parse + docx export tests."""

from io import BytesIO

from docx import Document

from src.protocol import parse_protocol, protocol_to_docx


def test_parse_and_docx_structure(sample_protocol: str) -> None:
    protocol = parse_protocol(sample_protocol)
    assert protocol.date == "18.03.2026"
    assert len(protocol.tasks) == 2

    document = Document(BytesIO(protocol_to_docx(protocol)))
    texts = [p.text for p in document.paragraphs]
    assert any("Протокол встречи" in t for t in texts)
    assert any(t == "Задачи" for t in texts)

    table = document.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Задача", "Ответственный", "Срок"]
    assert table.rows[1].cells[0].text == "Подготовить OpenAPI"


def test_docx_with_missing_fields(protocol_with_gaps: str) -> None:
    protocol = parse_protocol(protocol_with_gaps)
    document = Document(BytesIO(protocol_to_docx(protocol)))
    texts = [p.text for p in document.paragraphs]
    assert any("Дата: Не указан" in t for t in texts)
    assert document.tables[0].rows[1].cells[1].text == "Не указан"
