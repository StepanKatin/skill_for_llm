"""Parse fixed-format meeting protocol and export to .docx."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

PLACEHOLDER = "Не указан"
EMPTY_SECTION_PLACEHOLDER = "Нет данных"

SECTION_PATTERN = re.compile(
    r"^##\s+(Метаданные|Обсуждение|Решения|Задачи)\s*$",
    re.MULTILINE,
)
TASK_ROW_PATTERN = re.compile(
    r"^\|\s*(.*?)\s*\|\s*(.*?)\s*\|\s*(.*?)\s*\|\s*$",
    re.MULTILINE,
)


class ProtocolParseError(ValueError):
    pass


@dataclass(slots=True)
class TaskItem:
    title: str
    assignee: str = PLACEHOLDER
    due_date: str = PLACEHOLDER

    def normalized(self) -> TaskItem:
        return TaskItem(
            title=(self.title or "").strip() or PLACEHOLDER,
            assignee=(self.assignee or "").strip() or PLACEHOLDER,
            due_date=(self.due_date or "").strip() or PLACEHOLDER,
        )


@dataclass(slots=True)
class MeetingProtocol:
    title: str = "Протокол встречи"
    date: str = PLACEHOLDER
    participants: str = PLACEHOLDER
    discussion: str = EMPTY_SECTION_PLACEHOLDER
    decisions: str = EMPTY_SECTION_PLACEHOLDER
    tasks: list[TaskItem] = field(default_factory=list)

    def normalized(self) -> MeetingProtocol:
        return MeetingProtocol(
            title=(self.title or "").strip() or "Протокол встречи",
            date=(self.date or "").strip() or PLACEHOLDER,
            participants=(self.participants or "").strip() or PLACEHOLDER,
            discussion=(self.discussion or "").strip() or EMPTY_SECTION_PLACEHOLDER,
            decisions=(self.decisions or "").strip() or EMPTY_SECTION_PLACEHOLDER,
            tasks=[task.normalized() for task in self.tasks],
        )


def parse_protocol(text: str) -> MeetingProtocol:
    if not text or not text.strip():
        raise ProtocolParseError("protocol text is empty")

    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:markdown)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned).strip()

    title = "Протокол встречи"
    title_match = re.search(r"^#\s+(.+)$", cleaned, re.MULTILINE)
    if title_match:
        title = title_match.group(1).strip() or title

    sections = _split_sections(cleaned)
    meta = sections.get("Метаданные", "")
    return MeetingProtocol(
        title=title,
        date=_meta_value(meta, "Дата"),
        participants=_meta_value(meta, "Участники"),
        discussion=sections.get("Обсуждение", "").strip() or EMPTY_SECTION_PLACEHOLDER,
        decisions=sections.get("Решения", "").strip() or EMPTY_SECTION_PLACEHOLDER,
        tasks=_parse_tasks(sections.get("Задачи", "")),
    ).normalized()


def protocol_to_docx(protocol: MeetingProtocol) -> bytes:
    data = protocol.normalized()
    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)

    title = document.add_heading(data.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Метаданные", level=1)
    document.add_paragraph(f"Дата: {data.date}")
    document.add_paragraph(f"Участники: {data.participants}")

    document.add_heading("Обсуждение", level=1)
    document.add_paragraph(data.discussion or EMPTY_SECTION_PLACEHOLDER)

    document.add_heading("Решения", level=1)
    document.add_paragraph(data.decisions or EMPTY_SECTION_PLACEHOLDER)

    document.add_heading("Задачи", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Задача"
    header[1].text = "Ответственный"
    header[2].text = "Срок"

    tasks = data.tasks or []
    if not tasks:
        row = table.add_row().cells
        row[0].text = EMPTY_SECTION_PLACEHOLDER
        row[1].text = PLACEHOLDER
        row[2].text = PLACEHOLDER
    else:
        for task in tasks:
            row = table.add_row().cells
            row[0].text = task.title
            row[1].text = task.assignee
            row[2].text = task.due_date

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _split_sections(text: str) -> dict[str, str]:
    matches = list(SECTION_PATTERN.finditer(text))
    if not matches:
        raise ProtocolParseError(
            "required sections not found: Метаданные, Обсуждение, Решения, Задачи"
        )
    sections: dict[str, str] = {}
    for index, match in enumerate(matches):
        name = match.group(1)
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        sections[name] = text[start:end].strip()
    return sections


def _meta_value(meta_text: str, key: str) -> str:
    pattern = re.compile(
        rf"^[-*]\s*{re.escape(key)}\s*:\s*(.*)$",
        re.MULTILINE | re.IGNORECASE,
    )
    match = pattern.search(meta_text or "")
    if not match:
        return PLACEHOLDER
    return match.group(1).strip() or PLACEHOLDER


def _parse_tasks(tasks_raw: str) -> list[TaskItem]:
    rows: list[TaskItem] = []
    for match in TASK_ROW_PATTERN.finditer(tasks_raw or ""):
        cells = [cell.strip() for cell in match.groups()]
        if not any(cells):
            continue
        joined = "|".join(cells).lower()
        if set("".join(cells)) <= {"-", "|", ":", " "}:
            continue
        if cells[0].lower() == "задача" and cells[1].lower().startswith("ответствен"):
            continue
        if "---" in joined:
            continue
        rows.append(TaskItem(title=cells[0], assignee=cells[1], due_date=cells[2]))
    return rows
